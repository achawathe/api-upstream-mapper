"""Excel mapping report: one workbook, one sheet per question a reviewer asks."""

from collections import defaultdict
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from api_mapper.models import (
    Candidate,
    EvidenceEffect,
    FieldMapping,
    MappingRun,
    MappingStatus,
    TransformKind,
)
from api_mapper.reporting.comparison import explain_comparison

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILLS = {
    MappingStatus.AUTO: PatternFill("solid", fgColor="D9F2E3"),
    MappingStatus.MANUAL: PatternFill("solid", fgColor="D6E8FA"),
    MappingStatus.DERIVED: PatternFill("solid", fgColor="EDE7F6"),
    MappingStatus.REVIEW: PatternFill("solid", fgColor="FFF1CC"),
    MappingStatus.UNMAPPED: PatternFill("solid", fgColor="F9D9D6"),
    MappingStatus.REJECTED: PatternFill("solid", fgColor="E6E6E6"),
}
MAX_COLUMN_WIDTH = 60
WRAP = Alignment(wrap_text=True, vertical="top")


def write_excel_report(run: MappingRun, path: Path) -> Path:
    workbook = Workbook()
    _summary(workbook.active, run)
    _field_mapping(workbook.create_sheet("Field Mapping"), run.mappings)
    _comparisons(workbook.create_sheet("Candidate Comparison"), run.mappings)
    _needs_review(workbook.create_sheet("Needs Review"), run.mappings)
    _gaps(workbook.create_sheet("Gaps"), run.mappings)
    _transforms(workbook.create_sheet("Transforms"), run.mappings)
    _call_plan(workbook.create_sheet("Call Plan"), run)
    _errors(workbook.create_sheet("Error Codes"), run)
    _run_log(workbook.create_sheet("Run Log"), run)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _comparisons(sheet: Worksheet, mappings: list[FieldMapping]) -> None:
    rows = [
        (
            "Upstream field",
            "Chosen source",
            "Decision",
            "Supporting evidence",
            "Alternative",
            "Why not selected",
            "Uncertainties",
            "Conversions and conditions",
            "API references",
        )
    ]
    for mapping in mappings:
        comparison = explain_comparison(mapping)
        alternatives = comparison["alternatives"] or [{"source": "", "reason": ""}]
        for other in alternatives:
            rows.append(
                (
                    mapping.target.key,
                    _source_text(mapping.chosen),
                    comparison["summary"],
                    "; ".join(comparison["support"]),
                    other["source"],
                    other["reason"],
                    "; ".join(comparison["concerns"]),
                    "; ".join(comparison["requirements"]),
                    "\n".join(f"{r['side']} — {r['location']}: {r['text']}" for r in comparison["references"]),
                )
            )
    _table(sheet, rows)


def _summary(sheet: Worksheet, run: MappingRun) -> None:
    sheet.title = "Summary"
    sheet["A1"] = f"Mapping report: {run.upstream.title} {run.upstream.version}".strip()
    sheet["A1"].font = Font(bold=True, size=14)
    rows: list[Sequence[Any]] = [("Downstream API", "Title", "Source kind", "Endpoints")]
    rows += [(s.name, s.title, s.kind.value, len(s.endpoints)) for s in run.downstreams]
    _table(sheet, rows, start_row=3)

    counts = run.status_counts()
    start = len(rows) + 5
    _table(sheet, [("Status", "Fields"), *((status, count) for status, count in counts.items())], start_row=start)

    coverage_rows: list[Sequence[Any]] = [
        ("Upstream endpoint", "Fields", "Mapped", "Mapped %", "Required", "Required mapped", "Required %")
    ]
    by_endpoint: dict[str, list[FieldMapping]] = defaultdict(list)
    for mapping in run.mappings:
        by_endpoint[mapping.target.endpoint].append(mapping)
    for endpoint, mappings in by_endpoint.items():
        mapped = [m for m in mappings if m.is_mapped]
        required = [m for m in mappings if m.target.required]
        required_mapped = [m for m in required if m.is_mapped]
        coverage_rows.append(
            (
                endpoint,
                len(mappings),
                len(mapped),
                _ratio(mapped, mappings),
                len(required),
                len(required_mapped),
                _ratio(required_mapped, required),
            )
        )
    _table(sheet, coverage_rows, start_row=start + len(counts) + 3)


def _field_mapping(sheet: Worksheet, mappings: list[FieldMapping]) -> None:
    rows: list[Sequence[Any]] = [
        (
            "Upstream endpoint",
            "Variant",
            "Condition",
            "Upstream field",
            "Type",
            "Required",
            "Status",
            "Confidence",
            "Downstream API",
            "Downstream endpoint",
            "Downstream field",
            "Downstream type",
            "Transforms",
            "Why",
            "Notes",
            "Category",
        )
    ]
    for m in mappings:
        source = m.chosen.source if m.chosen else None
        rows.append(
            (
                m.target.endpoint,
                m.target.variant or "",
                m.target.condition or "",
                m.target.path,
                _type(m.target),
                "yes" if m.target.required else "",
                m.status.value,
                m.chosen.score if m.chosen else None,
                source.api if source else "",
                source.endpoint if source else "",
                source.path if source else "",
                _type(source) if source else "",
                _transforms_text(m.chosen),
                _why(m.chosen),
                "; ".join(m.notes),
                m.category,
            )
        )
    _table(sheet, rows)
    _color_by_status(sheet, mappings, status_column=7)


def _needs_review(sheet: Worksheet, mappings: list[FieldMapping]) -> None:
    rows: list[Sequence[Any]] = [
        (
            "Upstream endpoint",
            "Variant",
            "Upstream field",
            "Proposed source",
            "Score",
            "Why",
            "Notes",
            "Next best",
            "Next score",
        )
    ]
    for m in mappings:
        if m.status is not MappingStatus.REVIEW:
            continue
        runner_up = next((c for c in m.alternatives if c.vetoed_by is None), None)
        rows.append(
            (
                m.target.endpoint,
                m.target.variant or "",
                m.target.path,
                _source_text(m.chosen),
                m.chosen.score,
                _why(m.chosen),
                "; ".join(m.notes),
                _source_text(runner_up),
                runner_up.score if runner_up else None,
            )
        )
    _table(sheet, rows)


def _gaps(sheet: Worksheet, mappings: list[FieldMapping]) -> None:
    rows: list[Sequence[Any]] = [
        (
            "Upstream endpoint",
            "Variant",
            "Upstream field",
            "Required",
            "Status",
            "Closest candidate",
            "Score",
            "Why it was not used",
        )
    ]
    gaps = [m for m in mappings if m.status in (MappingStatus.UNMAPPED, MappingStatus.REJECTED)]
    gaps.sort(key=lambda m: (not m.target.required, m.target.endpoint, m.target.variant or "", m.target.path))
    for m in gaps:
        closest = m.alternatives[0] if m.alternatives else None
        rows.append(
            (
                m.target.endpoint,
                m.target.variant or "",
                m.target.path,
                "yes" if m.target.required else "",
                m.status.value,
                _source_text(closest),
                closest.score if closest else None,
                _rejection_reason(closest),
            )
        )
    _table(sheet, rows)


def _transforms(sheet: Worksheet, mappings: list[FieldMapping]) -> None:
    rows: list[Sequence[Any]] = [
        ("Upstream field", "Variant", "Source", "Transform", "Description", "Source value", "Upstream value")
    ]
    for m in mappings:
        if not m.chosen:
            continue
        for step in m.chosen.transforms:
            if step.kind is TransformKind.ENUM_MAP:
                for source_value, target_value in step.params.get("mapping", {}).items():
                    rows.append(
                        (
                            m.target.path,
                            m.target.variant or "",
                            _source_text(m.chosen),
                            step.kind.value,
                            "",
                            source_value,
                            target_value,
                        )
                    )
                for source_value in step.params.get("unmapped", []):
                    rows.append(
                        (
                            m.target.path,
                            m.target.variant or "",
                            _source_text(m.chosen),
                            step.kind.value,
                            "no match",
                            source_value,
                            "",
                        )
                    )
            else:
                rows.append(
                    (
                        m.target.path,
                        m.target.variant or "",
                        _source_text(m.chosen),
                        step.kind.value,
                        step.description,
                        "",
                        "",
                    )
                )
    _table(sheet, rows)


def _call_plan(sheet: Worksheet, run: MappingRun) -> None:
    rows: list[Sequence[Any]] = [
        (
            "Upstream endpoint",
            "Step",
            "Downstream API",
            "Downstream endpoint",
            "Only for variants",
            "Inputs",
            "Fan-out",
            "Fields supplied",
            "Pending review",
            "Purpose",
        )
    ]
    for plan in run.call_plans:
        if not plan.steps:
            rows.append((plan.upstream_endpoint, "", "", "", "", "no downstream source", "", 0, 0, ""))
        for step in plan.steps:
            inputs = "; ".join(
                f"{b.location} {b.parameter} <- {b.reference or b.source.value}" + (" (each item)" if b.fanout else "")
                for b in step.bindings
            )
            purpose = (
                "supplies inputs for later calls"
                if step.supplies_only_inputs
                else ("all fields await review" if step.awaits_review else "")
            )
            rows.append(
                (
                    plan.upstream_endpoint,
                    step.order,
                    step.api,
                    step.endpoint,
                    ", ".join(step.variants),
                    inputs,
                    "yes" if step.is_fanout else "",
                    step.supplies_fields,
                    step.pending_review_fields,
                    purpose,
                )
            )
        for unresolved in plan.unresolved_inputs:
            rows.append(
                (
                    plan.upstream_endpoint,
                    "",
                    "",
                    "",
                    "",
                    f"UNRESOLVED {unresolved}",
                    "",
                    "",
                    "",
                    "must come from caller context",
                )
            )
    _table(sheet, rows)


def _errors(sheet: Worksheet, run: MappingRun) -> None:
    rows: list[Sequence[Any]] = [
        (
            "Downstream API",
            "Endpoint",
            "Status",
            "Code",
            "Message",
            "Upstream status",
            "Upstream code",
            "Rule",
            "Confidence",
        )
    ]
    for e in run.error_mappings:
        d = e.downstream
        rows.append(
            (
                d.api,
                d.endpoint,
                d.http_status,
                d.code or "",
                d.message,
                e.upstream_status,
                e.upstream_code or "",
                e.rule,
                e.confidence,
            )
        )
    _table(sheet, rows)


def _run_log(sheet: Worksheet, run: MappingRun) -> None:
    rows: list[Sequence[Any]] = [("Stage", "State", "Message", "Counts", "Elapsed ms")]
    for event in run.events:
        counts = ", ".join(f"{k}={v}" for k, v in event.counts.items())
        rows.append((event.stage, event.state.value, event.message, counts, event.elapsed_ms or None))
    rows += [("warning", "", warning, "", None) for warning in run.warnings]
    _table(sheet, rows)


def _table(sheet: Worksheet, rows: Sequence[Sequence[Any]], start_row: int = 1) -> None:
    for offset, row in enumerate(rows):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=start_row + offset, column=column, value=value)
            cell.alignment = WRAP
            if offset == 0:
                cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    if start_row == 1 and len(rows) > 1:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(rows[0]))}{len(rows)}"
    _fit_columns(sheet, rows)


def _fit_columns(sheet: Worksheet, rows: Iterable[Sequence[Any]]) -> None:
    widths: dict[int, int] = {}
    for row in rows:
        for column, value in enumerate(row, start=1):
            widths[column] = max(widths.get(column, 8), min(MAX_COLUMN_WIDTH, len(str(value or "")) + 2))
    for column, width in widths.items():
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = max(sheet.column_dimensions[letter].width or 0, width)


def _color_by_status(sheet: Worksheet, mappings: list[FieldMapping], status_column: int) -> None:
    for row, mapping in enumerate(mappings, start=2):
        sheet.cell(row=row, column=status_column).fill = STATUS_FILLS[mapping.status]


def _type(node) -> str:
    return f"{node.type}({node.format})" if node.format else node.type


def _source_text(candidate: Candidate | None) -> str:
    if candidate is None:
        return ""
    return f"{candidate.source.api} {candidate.source.endpoint} {candidate.source.path}"


def _transforms_text(candidate: Candidate | None) -> str:
    return "; ".join(step.description for step in candidate.transforms) if candidate else ""


def _why(candidate: Candidate | None) -> str:
    if candidate is None:
        return ""
    strongest = sorted(
        (e for e in candidate.evidence if e.effect is EvidenceEffect.WEIGHTED and e.score >= 0.7),
        key=lambda e: -e.score * e.weight,
    )[:3]
    penalties = [e for e in candidate.evidence if e.effect is not EvidenceEffect.WEIGHTED]
    parts = [e.detail for e in strongest] + [f"penalty: {e.detail}" for e in penalties]
    return "; ".join(parts)


def _rejection_reason(candidate: Candidate | None) -> str:
    if candidate is None:
        return "no field in any downstream resembles this one"
    if candidate.vetoed_by:
        return f"vetoed: {candidate.vetoed_by}"
    penalties = [e.detail for e in candidate.evidence if e.effect is EvidenceEffect.MULTIPLIER]
    return "; ".join(penalties) or "score below the review threshold"


def _ratio(part: list, whole: list) -> str:
    return f"{len(part) / len(whole):.0%}" if whole else "n/a"
